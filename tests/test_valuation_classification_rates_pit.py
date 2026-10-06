"""Tests: classification rules, the discount rate and its Damodaran store, and point-in-time handling. Synthetic fixtures, no network."""
import hashlib
import json
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest

from valuation import classification as cl, discount_rate as dr
from valuation.common import CompanyPIT, DataNotPresent, Filed, PITViolation, SicInfo
from valuation.pit_companyfacts import CompanyFactsSource

ASOF = "2025-06-30"


def make_pit(sic, row=None, concepts=None, asof=ASOF, periods=("2024-12-31",)):
    series, reports = {}, []
    for pe in periods:
        filed = (date.fromisoformat(pe) + timedelta(days=60)).isoformat()
        for k, v in (row or {}).items():
            series.setdefault(k, {})[pe] = Filed(float(v), pe, filed, "A", "10-K", k)
        reports.append((pe, filed, "A", "10-K"))
    pit = CompanyPIT("SYN", asof, series=series, annual_reports=reports, sic=SicInfo(sic, "", "test") if sic else None)
    for name, v in (concepts or {}).items():
        pit.concepts[name] = {pe: Filed(float(v), pe, (date.fromisoformat(pe) + timedelta(days=60)).isoformat(), "A", "10-K", name)
                              for pe in periods}
    return pit


# ---- classification ----------------------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("sic,family", [("6331", "financial"), ("6022", "financial"), ("6211", "financial"), ("6199", "financial"),
                                        ("1311", "commodity"), ("1040", "commodity"), ("2911", "commodity"), ("3312", "commodity"),
                                        ("4412", "commodity"), ("2821", "commodity"), ("7372", "operating"), ("2834", "operating"),
                                        ("1389", "operating"), ("5311", "operating"), ("6282", "operating"), ("6411", "operating"),
                                        ("6798", "family_not_built"), ("6512", "family_not_built"), ("6719", "family_not_built"),
                                        ("6770", "family_not_built"), ("9995", "family_not_built")])
def test_sic_rules(sic, family):
    out = cl.classification("SYN", ASOF, make_pit(sic, {"total_assets": 100.0}))
    assert out["family"] == family
    assert out["reason"] and "SIC" in out["reason"]


def test_reit_and_conglomerate_are_annotated_family_not_built():
    r = cl.classification("SYN", ASOF, make_pit("6798"))
    assert r["family"] == "family_not_built" and "reit" in r["reason"]
    c = cl.classification("SYN", ASOF, make_pit("6719"))
    assert c["family"] == "family_not_built" and "conglomerate" in c["reason"]


def test_insurance_reserves_make_it_financial_and_the_balance_sheet_wins_over_sic():
    out = cl.classification("SYN", ASOF, make_pit("7389", {"insurance_reserves": 40.0, "total_assets": 100.0}))
    assert out["family"] == "financial"
    assert "balance sheet" in out["reason"] and "40.0%" in out["reason"]
    assert any("balance sheet says financial and wins" in c for c in out["conflicts"])


def test_insurance_reserves_immaterial_does_not_override_sic():
    out = cl.classification("SYN", ASOF, make_pit("7372", {"insurance_reserves": 2.0, "total_assets": 100.0}))
    assert out["family"] == "operating" and not out["conflicts"]


def test_lending_book_funded_by_borrowings_makes_it_financial():
    row = {"total_assets": 1000.0, "total_liabilities": 900.0, "debt_lt_noncurrent": 700.0}
    out = cl.classification("SYN", ASOF, make_pit("7389", row, {"LoansAndLeasesReceivableNetReportedAmount": 800.0}))
    assert out["family"] == "financial" and "lending" not in out["reason"].lower() or "loans" in out["reason"]
    assert out["tests"]["lending_test"] is True
    # a retailer with a card book that is a minority of assets is not financial
    small = cl.classification("SYN", ASOF, make_pit("5311", row, {"LoansAndLeasesReceivableNetReportedAmount": 200.0}))
    assert small["family"] == "operating" and small["tests"]["lending_test"] is False


def test_a_test_whose_inputs_are_not_filed_never_passes_by_default():
    out = cl.classification("SYN", ASOF, make_pit("7372", {"total_assets": 100.0}))
    assert out["tests"]["insurance_test"] is None and out["tests"]["lending_test"] is None
    assert len(out["tests"]["not_evaluated"]) == 2


def test_reit_by_sic_is_final_even_with_a_lending_book():
    row = {"total_assets": 1000.0, "total_liabilities": 900.0, "debt_lt_noncurrent": 700.0}
    out = cl.classification("SYN", ASOF, make_pit("6798", row, {"LoansAndLeasesReceivableNetReportedAmount": 800.0}))
    assert out["family"] == "family_not_built" and out["conflicts"]


def test_financial_sic_not_confirmed_by_balance_sheet_is_kept_and_annotated():
    out = cl.classification("SYN", ASOF, make_pit("6199", {"insurance_reserves": 1.0, "total_assets": 100.0}))
    assert out["family"] == "financial" and any("do not confirm" in c for c in out["conflicts"])


def test_no_annual_report_is_unclassifiable_and_no_sic_is_noted():
    pit = CompanyPIT("SYN", ASOF, sic=SicInfo("6331", "", "t"))
    assert cl.classification("SYN", ASOF, pit)["family"] is None
    out = cl.classification("SYN", ASOF, make_pit(None, {"total_assets": 10.0}))
    assert out["family"] == "operating" and any("no SIC" in n for n in out["notes"])


def test_the_sic_is_stamped_as_not_point_in_time():
    out = cl.classification("SYN", ASOF, make_pit("7372"))
    assert out["sic"]["point_in_time"] is False and any("current SEC assignment" in n for n in out["notes"])


# ---- point in time ------------------------------------------------------------------------------------------------------------------

def test_a_value_filed_on_the_valuation_date_is_a_violation_strictly_before_is_required():
    f = Filed(1.0, "2024-12-31", "2025-06-30", "A", "10-K", "revenue")
    with pytest.raises(PITViolation):
        CompanyPIT("SYN", "2025-06-30", series={"revenue": {"2024-12-31": f}})
    ok = Filed(1.0, "2024-12-31", "2025-06-29", "A", "10-K", "revenue")
    CompanyPIT("SYN", "2025-06-30", series={"revenue": {"2024-12-31": ok}})
    with pytest.raises(PITViolation):
        CompanyPIT("SYN", "2025-06-30", annual_reports=[("2024-12-31", "2025-07-01", "A", "10-K")])


def test_classification_checks_the_asof_it_is_given():
    with pytest.raises(ValueError):
        cl.classification("SYN", "2024-01-01", make_pit("7372"))


def _cf_row(val, end, start, accn, filed, form="10-K", fy=2024):
    r = {"end": end, "val": val, "accn": accn, "fy": fy, "fp": "FY", "form": form, "filed": filed}
    if start:
        r["start"] = start
    return r


def test_companyfacts_adapter_first_filed_own_period_strictly_before(tmp_path):
    rev = [_cf_row(100, "2022-12-31", "2022-01-01", "K22", "2023-03-01", fy=2022),
           _cf_row(110, "2023-12-31", "2023-01-01", "K23", "2024-03-01", fy=2023),
           _cf_row(105, "2022-12-31", "2022-01-01", "K23", "2024-03-01", fy=2023),         # a comparative, restated: ignored
           _cf_row(112, "2023-12-31", "2023-01-01", "K23A", "2024-09-01", form="10-K/A", fy=2023),   # an amendment: not first filed
           _cf_row(120, "2024-12-31", "2024-01-01", "K24", "2025-03-01", fy=2024),
           _cf_row(130, "2025-12-31", "2025-01-01", "K25", "2026-03-01", fy=2025)]          # filed after the valuation date
    assets = [_cf_row(500, "2022-12-31", None, "K22", "2023-03-01", fy=2022), _cf_row(520, "2023-12-31", None, "K23", "2024-03-01"),
              _cf_row(540, "2024-12-31", None, "K24", "2025-03-01")]
    facts = {"cik": 1, "facts": {"us-gaap": {"Revenues": {"units": {"USD": rev}}, "Assets": {"units": {"USD": assets}}},
                                 "dei": {"EntityCommonStockSharesOutstanding": {"units": {"shares": [
                                     {"end": "2025-02-20", "val": 50, "accn": "K24", "form": "10-K", "filed": "2025-03-01"},
                                     {"end": "2026-02-20", "val": 99, "accn": "K25", "form": "10-K", "filed": "2026-03-01"}]}}}}}
    d = tmp_path / "facts"
    d.mkdir()
    (d / "CIK0000000001.json").write_text(json.dumps(facts))
    sd = tmp_path / "subs"
    sd.mkdir()
    (sd / "CIK0000000001.json").write_text(json.dumps({"sic": "7372", "sicDescription": "Services-Prepackaged Software"}))
    src = CompanyFactsSource(facts_dir=d, cik_map={"SYN": "1"}, submissions_dirs=[sd])
    pit = src.load("SYN", "2025-06-30")
    assert pit.periods() == ["2022-12-31", "2023-12-31", "2024-12-31"]
    assert pit.series["revenue"]["2023-12-31"].value == 110 and pit.series["revenue"]["2023-12-31"].accession == "K23"
    assert pit.series["revenue"]["2022-12-31"].value == 100               # its own report, not the restated comparative
    assert "2025-12-31" not in pit.series["revenue"]                       # filed after asof
    assert pit.series["total_assets"]["2024-12-31"].value == 540
    assert pit.shares_outstanding.value == 50                               # the 2026 cover count is not visible
    assert pit.sic.code == "7372" and pit.sic.point_in_time is False
    # the same company one day after the K24 filing date still sees K24; on the filing date itself it does not
    assert src.load("SYN", "2025-03-02").periods()[-1] == "2024-12-31"
    assert src.load("SYN", "2025-03-01").periods()[-1] == "2023-12-31"


# ---- discount rate -------------------------------------------------------------------------------------------------------------------

def test_hamada_and_cost_of_equity_and_wacc_by_hand():
    bl = dr.relever_beta(1.0, 0.25, 0.5)
    assert bl == pytest.approx(1.375)
    coe = dr.cost_of_equity(0.04, bl, 0.05)
    assert coe == pytest.approx(0.10875)
    assert dr.wacc(coe, 0.06, 0.25, 0.5) == pytest.approx((1 / 1.5) * 0.10875 + (0.5 / 1.5) * 0.06 * 0.75)
    assert dr.wacc(coe, None, 0.25, 0.0) == coe                  # no debt: no cost of debt needed
    with pytest.raises(DataNotPresent):
        dr.wacc(coe, None, 0.25, 0.5)


def test_compute_discount_rate_assembles_and_flags_book_de_and_placeholders():
    p = dr.PlaceholderRates(True, rf=0.04, erp=0.05, beta_unlevered=1.0)
    d = dr.compute_discount_rate("2025-06-30", tax=0.25, de_ratio=0.5, de_basis="book", kd_pretax=0.06, industry="x", rf_provider=p.rf,
                                 beta_provider=p.beta, erp_provider=p.erp, placeholder=True)
    assert d.coe == pytest.approx(0.10875) and d.placeholder and any("book" in w for w in d.warnings) and any("PLACEHOLDER" in w for w in d.warnings)


def test_discount_rate_for_state_refuses_unmeasurable_inputs():
    p = dr.PlaceholderRates(True)
    kw = dict(industry="x", rf_provider=p.rf, beta_provider=p.beta, erp_provider=p.erp, placeholder=True)
    with pytest.raises(DataNotPresent):
        dr.discount_rate_for_state({"tax_rate": 0.2, "de_book": None, "kd_effective": 0.05}, "2025-06-30", **kw)
    with pytest.raises(DataNotPresent):
        dr.discount_rate_for_state({"tax_rate": 0.2, "de_book": 0.4, "kd_effective": None}, "2025-06-30", **kw)
    ok = dr.discount_rate_for_state({"tax_rate": 0.2, "de_book": 0.0, "kd_effective": None}, "2025-06-30", **kw)
    assert ok.wacc == pytest.approx(ok.coe)


BETAS_CSV = """Date updated:,{date},,,,,,
Created by:,x,,,,,,
Industry Name,Number of firms,Beta ,D/E Ratio,Effective Tax rate,Unlevered beta,Cash/Firm value,Unlevered beta corrected for cash
Steel,19,1.06,23.51%,9.28%,0.90,4.38%,{beta}
Software,309,1.28,5.58%,5.51%,1.23,1.83%,NA
"""
HIST_CSV = """Date updated:,4-Jan-20,,
,,,
Year,Earnings Yield,T.Bond Rate,Implied ERP (FCFE),Implied ERP with risk adjusted riskfree rate
2019,5.0%,2.0%,5.20%,
2020,4.0%,1.0%,4.72%,
2021,3.0%,1.5%,4.24%,
"""


def make_store(tmp_path, snapshots=(("5-Jan-20", "0.80"), ("5-Jan-22", "0.95"))):
    root = tmp_path / "dam"
    (root / "csv").mkdir(parents=True)
    files = {}

    def add(name, csv_name, text):
        (root / name).write_bytes(text.encode("utf-8"))                  # stands in for the .xls: only its hash matters
        (root / "csv" / csv_name).write_text(text, encoding="utf-8")
        files[name] = {"sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(), "bytes": len(text)}
    for i, (dte, beta) in enumerate(snapshots):
        add(f"archives_betas{i}.xls", f"archives_betas{i}__2_Industry_Averages.csv", BETAS_CSV.format(date=dte, beta=beta))
    add("datasets_histimpl.xls", "datasets_histimpl__1_Historical_Impl_Premiums.csv", HIST_CSV)
    (root / "MANIFEST.json").write_text(json.dumps({"files": files}))
    return root


def test_store_picks_the_newest_beta_file_strictly_before_the_valuation_date_and_records_hashes(tmp_path):
    s = dr.DamodaranStore(make_store(tmp_path))
    prov = s.beta_provider()
    v, label, meta = prov("2021-06-01", "Steel")
    assert v == 0.80 and meta["snapshot_date"] == "2020-01-05" and len(meta["file_sha256"]) == 64 and len(meta["csv_sha256"]) == 64
    assert prov("2022-01-06", "Steel")[0] == 0.95
    assert prov("2022-01-05", "Steel")[0] == 0.80      # the file dated the same day is not used: strictly before
    with pytest.raises(DataNotPresent):
        prov("2020-01-05", "Steel")                    # no file strictly before
    with pytest.raises(DataNotPresent):
        prov("2021-06-01", "Unknown industry")
    with pytest.raises(DataNotPresent):
        prov("2021-06-01", "Software")                  # NA unlevered beta: the unlevered provider refuses (the levered one has a value, see the next test)


def test_the_levered_provider_returns_the_betas_sheet_beta_column_from_the_same_point_in_time_file(tmp_path):
    """Addendum I3: the column "Beta " (the levered industry beta) of the newest file dated strictly before the valuation date; hashes recorded; no substitute."""
    s = dr.DamodaranStore(make_store(tmp_path))
    prov = s.levered_beta_provider()
    v, label, meta = prov("2021-06-01", "Steel")
    assert v == 1.06 and meta["snapshot_date"] == "2020-01-05" and meta["column"].startswith("Beta") and len(meta["file_sha256"]) == 64
    assert prov("2021-06-01", "Software")[0] == 1.28    # an industry with an NA unlevered beta still has its levered beta
    assert prov("2022-01-05", "Steel")[0] == 1.06       # strictly before: the file dated the same day is not used
    with pytest.raises(DataNotPresent):
        prov("2020-01-05", "Steel")
    with pytest.raises(DataNotPresent):
        prov("2021-06-01", "Unknown industry")
    d = dr.compute_discount_rate("2021-06-01", tax=0.25, de_ratio=3.0, de_basis="book", kd_pretax=0.05, industry="Steel", rf_provider=lambda a: (0.02, "rf"),
                                 beta_provider=prov, erp_provider=s.erp_provider(), levered_beta=True)
    assert d.beta_levered == 1.06 and d.beta_unlevered is None and d.coe == pytest.approx(0.02 + 1.06 * 0.0472)     # no Hamada: D/E = 3 does not enter


def test_store_erp_uses_the_year_end_row_after_the_publication_lag(tmp_path):
    s = dr.DamodaranStore(make_store(tmp_path))
    e = s.erp_provider()
    assert e("2021-06-01")[0] == pytest.approx(0.0472)                    # year-end 2020 is known by mid-2021
    assert e("2021-01-04")[0] == pytest.approx(0.0520)                    # 2020 year-end + 5 days has not passed: 2019
    assert e("2021-01-07")[0] == pytest.approx(0.0472)
    with pytest.raises(DataNotPresent):
        e("2019-12-31")                                                   # no earlier row


def add_monthly(root, rows, name=dr.ERP_MONTHLY_FILE, sheet=dr.ERP_MONTHLY_SHEET):
    """Adds a stand-in monthly ERP workbook (openpyxl) to a store made by make_store and lists it in the manifest with its real sha256."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["Start of month", "S&P 500", "T.Bond Rate", "ERP (T12 m with sustainable payout)", "ERP (T12m)"])
    for r in rows:
        ws.append(r)
    wb.create_sheet("Other")
    wb.save(root / name)
    man = json.loads((root / "MANIFEST.json").read_text())
    man["files"][name] = {"sha256": hashlib.sha256((root / name).read_bytes()).hexdigest(), "bytes": (root / name).stat().st_size}
    (root / "MANIFEST.json").write_text(json.dumps(man))


MONTHLY_ROWS = [[datetime(2015, 12, 1), 2000, 0.02, None, 0.0600],                 # the sustainable-payout column is empty before 2016, as in the real file
                [datetime(2016, 1, 1), 1900, 0.02, 0.0516, 0.0600],
                [datetime(2020, 2, 1), 3200, 0.015, 0.0522, 0.0524],
                [datetime(2020, 3, 1), 2950, 0.011, 0.0577, 0.0569],
                [datetime(2020, 4, 1), 2585, 0.007, 0.0652, 0.0616]]


def test_monthly_erp_takes_the_newest_start_of_month_row_on_or_before_the_date_and_never_a_later_one(tmp_path):
    """Addendum K: for a month-end date T the row of that month's first day; the next month's row is invisible; a row dated T itself counts (on or before)."""
    root = make_store(tmp_path)
    add_monthly(root, MONTHLY_ROWS)
    e = dr.DamodaranStore(root).erp_provider()
    v, label, meta = e("2020-03-31")
    assert v == pytest.approx(0.0577) and meta["row_date"] == "2020-03-01" and meta["source"] == "monthly" and "start of month 2020-03-01" in label and len(meta["file_sha256"]) == 64
    assert e("2020-04-01")[0] == pytest.approx(0.0652)               # a row dated the date itself is on or before it
    assert e("2020-03-01")[0] == pytest.approx(0.0577)
    assert e("2020-02-29")[0] == pytest.approx(0.0522)
    assert e("2020-03-31")[2]["row_date"] != "2020-04-01"            # no look-ahead
    with pytest.raises(DataNotPresent, match="stale"):
        e("2020-07-31")                                                # the newest row is more than 62 days old: a hole in a monthly file is a refusal, not a stale value


def test_monthly_erp_falls_back_to_the_annual_year_end_row_before_the_columns_first_value_and_says_so(tmp_path):
    root = make_store(tmp_path)
    # a store whose first value is 2020-03-01 (the 2015-12 row holds None, not a value): a date before it uses the annual year-end row, labelled
    add_monthly(root, MONTHLY_ROWS[:1] + [[datetime(2020, 3, 1), 2950, 0.011, 0.0577, 0.0569]])
    e2 = dr.DamodaranStore(root).erp_provider()
    v, label, meta = e2("2020-01-31")
    assert meta["source"] == "annual" and "annual fallback" in label and meta["fallback_reason"] and v == pytest.approx(0.0520)       # year-end 2019 (5.20%)


def test_monthly_erp_without_a_monthly_file_in_the_manifest_uses_the_annual_path_and_a_listed_file_is_verified(tmp_path):
    s = dr.DamodaranStore(make_store(tmp_path))
    assert s.erp_provider()("2021-06-01")[2]["source"] == "annual"
    root = make_store(tmp_path / "m")
    add_monthly(root, MONTHLY_ROWS)
    (root / dr.ERP_MONTHLY_FILE).write_bytes(b"tampered")
    with pytest.raises(DataNotPresent, match="hash mismatch"):
        dr.DamodaranStore(root).erp_provider()("2020-03-31")
    root2 = make_store(tmp_path / "n")
    add_monthly(root2, MONTHLY_ROWS)
    (root2 / dr.ERP_MONTHLY_FILE).unlink()
    with pytest.raises(DataNotPresent, match="not in the manifest or not on disk"):
        dr.DamodaranStore(root2).erp_provider()("2020-03-31")


def test_monthly_erp_parser_skips_empty_cells_refuses_a_percent_and_a_missing_column():
    hdr = ["Start of month", "S&P 500", "ERP (T12 m with sustainable payout)"]
    assert dr.parse_erp_monthly_rows([hdr, [datetime(2016, 1, 1), 1, None], [datetime(2016, 2, 1), 1, "NA"], [datetime(2016, 3, 1), 1, 0.05], ["note"], [datetime(2016, 4, 1), 1, "5.10%"]])         == [(date(2016, 3, 1), 0.05), (date(2016, 4, 1), 0.051)]
    with pytest.raises(DataNotPresent):
        dr.parse_erp_monthly_rows([hdr, [datetime(2016, 1, 1), 1, 5.1]])         # a percent, not a fraction: refused, not rescaled
    with pytest.raises(DataNotPresent):
        dr.parse_erp_monthly_rows([["Start of month", "x"], [datetime(2016, 1, 1), 1]])
    with pytest.raises(DataNotPresent):
        dr.parse_erp_monthly_rows([["Month", "x"]])


def test_store_refuses_a_hash_mismatch_and_a_missing_directory(tmp_path):
    root = make_store(tmp_path)
    (root / "archives_betas0.xls").write_bytes(b"tampered")
    with pytest.raises(DataNotPresent, match="hash mismatch"):
        dr.DamodaranStore(root).beta_provider()("2021-06-01", "Steel")
    with pytest.raises(DataNotPresent, match="not present"):
        dr.DamodaranStore(tmp_path / "nowhere")
    with pytest.raises(DataNotPresent, match="indname"):
        dr.DamodaranStore(make_store(tmp_path / "b")).industry_for_sic("1311")


def test_parsers_on_documented_layouts():
    rows = [["Industry Name", "Number of firms", "Beta ", "D/E Ratio", "Tax rate", "Unlevered beta", "Cash/Firm value",
             "Unlevered beta corrected for cash"], ["Auto Parts", "35", "1.34", "41.46%", "15.00%", "1.02", "9.45%", "1.13"]]
    t = dr.parse_betas_rows(rows)
    assert t["Auto Parts"] == {"beta_unlevered": 1.13, "beta_levered": 1.34, "de_ratio": 0.4146, "tax": 0.15, "firms": 35.0}
    with pytest.raises(DataNotPresent):
        dr.parse_betas_rows([["x"]])
    assert dr.parse_updated_cell([["Date updated:", "5-Jan-26"]]) == date(2026, 1, 5)
    with pytest.raises(DataNotPresent):
        dr.parse_histimpl_rows([["Year", "Implied ERP (FCFE)"], ["2020", "47.2"]])       # a percent, not a fraction: refused, not rescaled


REAL = Path(__file__).resolve().parents[3] / "data" / "damodaran"


def _real_root():
    try:
        return dr.default_root()
    except DataNotPresent:
        return None


@pytest.mark.skipif(_real_root() is None, reason="the approved Damodaran files are not on this machine")
def test_the_approved_damodaran_files_parse_select_by_date_and_record_hashes():
    s = dr.DamodaranStore(_real_root())
    snaps = s.beta_snapshots()
    assert [e["date"].year for e in snaps] == list(range(2017, 2027))
    for e in snaps:
        table = dr.parse_betas_rows(e["_rows"])
        assert len(table) > 80 and "Steel" in table and 0.0 < table["Steel"]["beta_unlevered"] < 3.0
        assert 0.0 < table["Steel"]["beta_levered"] < 3.0 and 0.0 < table["Banks (Regional)"]["beta_levered"] < 3.0        # Addendum I3: the "Beta" column, in every file
    v, label, meta = s.beta_provider()("2021-06-01", "Steel")
    assert meta["snapshot_date"] == "2021-01-05" and len(meta["file_sha256"]) == 64
    erp, label, meta = s.erp_provider()("2021-06-01")
    assert erp == pytest.approx(0.0420) and meta["source"] == "monthly" and meta["row_date"] == "2021-06-01" and len(meta["file_sha256"]) == 64        # Addendum K: the monthly row
    erp, label, meta = s.erp_provider(monthly=False)("2021-06-01")
    assert 0.03 < erp < 0.08 and meta["year_end"] == 2020 and meta["source"] == "annual"                                                            # the annual path is unchanged


# ---- Addendum H4: family_not_built for utilities and pipelines, and for pre-profit companies --------------------------------------------------------

@pytest.mark.parametrize("sic", ["4900", "4911", "4922", "4923", "4924", "4925", "4931", "4941", "4949"])
def test_regulated_utilities_and_pipelines_are_family_not_built(sic):
    out = cl.classification("SYN", ASOF, make_pit(sic, {"total_assets": 100.0}))
    assert out["family"] == "family_not_built" and "regulated_utility_or_pipeline" in out["reason"] and "family not built" in out["reason"]


@pytest.mark.parametrize("sic,family", [("4899", "operating"), ("4950", "operating"), ("4813", "operating"), ("4610", "operating")])
def test_the_utility_range_is_exactly_4900_to_4949(sic, family):
    assert cl.classification("SYN", ASOF, make_pit(sic, {"total_assets": 100.0}))["family"] == family


def _oi_pit(sic, ois):
    periods = tuple(f"{2025 - len(ois) + i}-12-31" for i in range(len(ois)))
    pit = make_pit(sic, {"total_assets": 100.0}, periods=periods)
    for pe, v in zip(periods, ois):
        if v is not None:
            filed = (date.fromisoformat(pe) + timedelta(days=60)).isoformat()
            pit.series.setdefault("operating_income", {})[pe] = Filed(float(v), pe, filed, "A", "10-K", "operating_income")
    return pit


def test_pre_profit_is_operating_income_at_or_below_zero_in_three_of_the_last_four_years():
    # the last four years known: -5, -1, 0, +9  -> three at or below zero (0.0 is a value, and a loss): pre-profit
    out = cl.classification("SYN", ASOF, _oi_pit("2834", [20, -5, -1, 0, 9]))
    assert out["family"] == "family_not_built" and out["reason"].startswith("pre_profit: operating income <= 0 in 3 of the last 4")
    assert out["tests"]["pre_profit"]["pre_profit"] is True and any("operating income was <= 0" in c for c in out["conflicts"])
    # two of four is not
    ok = cl.classification("SYN", ASOF, _oi_pit("2834", [-5, -1, 4, 9]))
    assert ok["family"] == "operating" and ok["tests"]["pre_profit"]["loss_years"] == 2
    # only the last four count: an old loss year outside the window does not
    assert cl.classification("SYN", ASOF, _oi_pit("2834", [-9, -9, -9, 5, 6, 7, -1]))["family"] == "operating"
    # a company with three fiscal years, all losses (a young company): pre-profit
    assert cl.classification("SYN", ASOF, _oi_pit("7372", [-3, -4, -2]))["family"] == "family_not_built"
    # a year with no operating income filed is not counted either way
    assert cl.classification("SYN", ASOF, _oi_pit("7372", [-3, None, -2, 5]))["family"] == "operating"
    assert cl.classification("SYN", ASOF, _oi_pit("7372", [-3, None, -2, -1]))["family"] == "family_not_built"


def test_pre_profit_applies_to_commodity_but_not_to_financials():
    assert cl.classification("SYN", ASOF, _oi_pit("1311", [-1, -2, -3, -4]))["family"] == "family_not_built"
    assert cl.classification("SYN", ASOF, _oi_pit("6022", [-1, -2, -3, -4]))["family"] == "financial"
