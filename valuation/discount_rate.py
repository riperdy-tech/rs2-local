"""discount_rate.py - SPEC 3.4: cost of equity = risk-free + relevered beta x equity risk premium.

  relevered beta = industry unlevered beta x (1 + (1 - tax) x D/E)         (Hamada, with the company's OWN D/E)
  cost of equity = rf + relevered beta x ERP
  WACC           = E/V x cost of equity + D/V x pre-tax cost of debt x (1 - tax),   E/V = 1/(1+D/E), D/V = (D/E)/(1+D/E)

Data sources are pluggable (three small provider callables, each returning its value, a label and a data-set fingerprint):
  rf   : rf(asof)                    -> (fraction, label)
  beta : beta(asof, industry)        -> (unlevered beta, label, {"file_sha256", "snapshot_date", ...})
  erp  : erp(asof)                   -> (fraction, label, {...})
`DamodaranStore` implements beta and erp on the operator-approved Damodaran files (stocks-workspace/data/damodaran: the original .xls,
MANIFEST.json with each file's sha256, and every sheet converted to csv/ by Excel). NOTHING IS DOWNLOADED by this code. Each load
re-hashes the .xls against the manifest; every result records the .xls sha256 and the sha256 of the csv it read.

Point in time (no look-ahead, nothing substituted from a later file):
  * beta: the newest betas file whose own "Date updated:" cell is strictly before `asof` (archives betasYY = early January YY+1;
    datasets_betas = 5-Jan-26). Industry unlevered beta = "Unlevered beta corrected for cash" (Damodaran's pure-play beta; he unlevers
    with a marginal tax rate stated in the sheet).
  * ERP: histimpl.xls column "Implied ERP (FCFE)", row of the newest year-end Y with Dec 31 of Y + publication_lag_days (5) before
    `asof`: a year-end value is used only after early January of the next year, the date his betas files carry. (The sheet's own
    "Date updated" cell is stale and is not used.)
The SIC -> Damodaran industry mapping (indname.xls) is NOT among the approved files; the industry is passed explicitly until it is.

Concentration risk is carried in scenario cash flows, never as an extra premium (SPEC 3.4). The closure premium is a diagnostic, not here.
"""
from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path
from typing import Callable, Dict, List, Optional

from .common import DataNotPresent, InputError, parse_date, sha256_file

PUBLICATION_LAG_DAYS = 5
BETA_COLUMN = "unlevered beta corrected for cash"
ERP_COLUMN = "implied erp (fcfe)"
APPROVAL_NOTE = ("Damodaran data is not present at the expected location (approved 2026-10-06; see stocks-workspace/data/damodaran/"
                 "MANIFEST.json).")


# ---- the arithmetic ------------------------------------------------------------------------------------------------------------

def relever_beta(beta_unlevered: float, tax: float, de_ratio: float) -> float:
    """Hamada: beta_L = beta_U x (1 + (1 - tax) x D/E)."""
    if de_ratio < 0:
        raise InputError("D/E must not be negative")
    return beta_unlevered * (1.0 + (1.0 - tax) * de_ratio)


def cost_of_equity(rf: float, beta_levered: float, erp: float) -> float:
    return rf + beta_levered * erp


def wacc(coe: float, kd_pretax: Optional[float], tax: float, de_ratio: float) -> float:
    """Weights from D/E: E/V = 1/(1+D/E). With no debt the pre-tax cost of debt is not needed (None allowed)."""
    ev = 1.0 / (1.0 + de_ratio)
    dv = de_ratio / (1.0 + de_ratio)
    if dv == 0.0:
        return coe
    if kd_pretax is None:
        raise DataNotPresent("the company has debt but its pre-tax cost of debt cannot be measured from its filings")
    return ev * coe + dv * kd_pretax * (1.0 - tax)


@dataclass
class DiscountRate:
    asof: str
    rf: float
    erp: float
    beta_unlevered: float
    de_ratio: float
    de_basis: str                   # "book" (filed values, price-free) or "market"
    tax: float
    beta_levered: float
    coe: float
    kd_pretax: Optional[float]
    wacc: float
    sources: Dict[str, object] = field(default_factory=dict)
    placeholder: bool = False
    warnings: List[str] = field(default_factory=list)

    def as_dict(self):
        return dict(self.__dict__)


def compute_discount_rate(asof, *, tax: float, de_ratio: float, de_basis: str, kd_pretax: Optional[float],
                          industry: str, rf_provider: Callable, beta_provider: Callable, erp_provider: Callable,
                          placeholder: bool = False) -> DiscountRate:
    """Assemble the discount rate for `asof` from the three providers. Any provider may raise DataNotPresent: the error propagates
    (a refusal, never a fallback)."""
    asof_d = parse_date(asof)
    rf, rf_label = rf_provider(asof_d)
    bu, b_label, b_meta = beta_provider(asof_d, industry)
    erp, e_label, e_meta = erp_provider(asof_d)
    bl = relever_beta(bu, tax, de_ratio)
    coe = cost_of_equity(rf, bl, erp)
    w = wacc(coe, kd_pretax, tax, de_ratio)
    warns = []
    if de_basis == "book":
        warns.append("D/E is the book (filed) ratio, not the market ratio: Hamada's relevering is defined on market D/E")
    if placeholder:
        warns.append("PLACEHOLDER rates: for tests and plumbing only; never a valuation input")
    return DiscountRate(asof=asof_d.isoformat(), rf=rf, erp=erp, beta_unlevered=bu, de_ratio=de_ratio, de_basis=de_basis, tax=tax,
                        beta_levered=bl, coe=coe, kd_pretax=kd_pretax, wacc=w,
                        sources={"rf": rf_label, "beta": b_label, "beta_meta": b_meta, "erp": e_label, "erp_meta": e_meta,
                                 "industry": industry},
                        placeholder=placeholder, warnings=warns)


def discount_rate_for_state(st: dict, asof, *, industry: str, rf_provider: Callable, beta_provider: Callable, erp_provider: Callable,
                            de_ratio: Optional[float] = None, de_basis: str = "book", placeholder: bool = False) -> DiscountRate:
    """The discount rate of a company from its common.build_state(): its own tax rate, its own D/E (book, from the filings, unless
    `de_ratio` is given with its basis) and its own measured pre-tax cost of debt (interest expense / average debt). A D/E or cost of
    debt that cannot be measured is a refusal (DataNotPresent), not a default."""
    if de_ratio is None:
        de_ratio = st.get("de_book")
        de_basis = "book"
    if de_ratio is None:
        raise DataNotPresent("D/E cannot be measured from the filings (debt or parent equity missing, or equity not positive)")
    return compute_discount_rate(asof, tax=st["tax_rate"], de_ratio=de_ratio, de_basis=de_basis, kd_pretax=st.get("kd_effective"),
                                 industry=industry, rf_provider=rf_provider, beta_provider=beta_provider, erp_provider=erp_provider,
                                 placeholder=placeholder)


# ---- readers (rows as lists of cells) and parsers -------------------------------------------------------------------------------------

def read_rows(path, sheet: Optional[str] = None) -> List[list]:
    """The rows of a .csv, .xlsx (openpyxl) or .xls (xlrd) file. The Damodaran sheets were converted to CSV by Excel (xlrd is not
    installed and must not be installed without the operator's approval), so the CSVs are what the store reads."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".csv":
        raw = p.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1252")          # Excel's csv export on Windows
        return [row for row in csv.reader(text.splitlines())]
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if ext == ".xls":
        try:
            import xlrd
        except ImportError:
            raise DataNotPresent("reading a .xls needs xlrd, which is not installed; read the csv conversion instead")
        wb = xlrd.open_workbook(str(p))
        ws = wb.sheet_by_name(sheet) if sheet else wb.sheet_by_index(0)
        return [ws.row_values(i) for i in range(ws.nrows)]
    raise InputError(f"unsupported spreadsheet type {ext}")


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s if s is not None else "")).strip().lower()


def _num(x):
    """A cell as a float: '40.20%' -> 0.402, '1.21' -> 1.21, 'NA' / '#DIV/0!' / '' -> None."""
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return None if x != x else float(x)
    t = str(x).strip().replace(",", "")
    if not t or t.upper() in ("NA", "N/A") or t.startswith("#"):
        return None
    try:
        if t.endswith("%"):
            return float(t[:-1]) / 100.0
        return float(t)
    except ValueError:
        return None


_MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def parse_updated_cell(rows) -> date:
    """The 'Date updated:' cell of a Damodaran sheet ('5-Jan-26'): the date the file's data describes."""
    for r in rows[:5]:
        if r and _norm(r[0]).startswith("date updated"):
            cell = r[1] if len(r) > 1 else ""
            if isinstance(cell, date):
                return cell
            m = re.match(r"^\s*(\d{1,2})-([A-Za-z]{3})[a-z]*-(\d{2,4})\s*$", str(cell))
            if m:
                yy = int(m.group(3))
                yy = yy + 2000 if yy < 100 else yy
                return date(yy, _MONTHS[m.group(2).lower()], int(m.group(1)))
    raise DataNotPresent("no parsable 'Date updated:' cell in the first rows of the sheet")


def parse_betas_rows(rows, beta_column: str = BETA_COLUMN) -> Dict[str, dict]:
    """{industry name: {"beta_unlevered", "de_ratio", "tax", "firms"}} from the rows of the industry-averages sheet."""
    hdr = None
    for i, r in enumerate(rows):
        if r and _norm(r[0]) == "industry name":
            hdr = i
            break
    if hdr is None:
        raise DataNotPresent("betas sheet: no header row starting with 'Industry Name'")
    cols = {}
    for j, c in enumerate(rows[hdr]):
        cols.setdefault(_norm(c), j)
    if beta_column not in cols:
        raise DataNotPresent(f"betas sheet: no column {beta_column!r} (found {sorted(cols)})")
    tax_col = cols.get("effective tax rate", cols.get("tax rate"))
    out = {}
    for r in rows[hdr + 1:]:
        if not r or not str(r[0]).strip():
            continue
        j = cols[beta_column]
        bu = _num(r[j]) if j < len(r) else None
        if bu is None:
            continue

        def at(j):
            return _num(r[j]) if j is not None and j < len(r) else None
        out[str(r[0]).strip()] = {"beta_unlevered": bu, "de_ratio": at(cols.get("d/e ratio")),
                                  "tax": at(tax_col), "firms": at(cols.get("number of firms"))}
    if not out:
        raise DataNotPresent("betas sheet: no industry rows parsed")
    return out


def parse_histimpl_rows(rows, column: str = ERP_COLUMN) -> Dict[int, float]:
    """{year: implied ERP (fraction)} from the rows of the implied-premium history sheet; the row of year Y is the year-end-Y value."""
    hdr = None
    for i, r in enumerate(rows):
        if r and _norm(r[0]) == "year":
            hdr = i
            break
    if hdr is None:
        raise DataNotPresent("histimpl sheet: no header row starting with 'Year'")
    j = next((k for k, c in enumerate(rows[hdr]) if _norm(c) == column), None)
    if j is None:
        raise DataNotPresent(f"histimpl sheet: no column {column!r}")
    out = {}
    for r in rows[hdr + 1:]:
        if not r:
            continue
        y = _num(r[0])
        if y is None or y != int(y) or not (1900 < y < 2200):
            continue
        v = _num(r[j]) if j < len(r) else None
        if v is None:
            continue
        if not (0.0 < v < 0.5):
            raise DataNotPresent(f"histimpl sheet: implied ERP {v} for {int(y)} is not a plausible fraction - refusing to use it")
        out[int(y)] = v
    if not out:
        raise DataNotPresent("histimpl sheet: no ERP rows parsed")
    return out


def parse_indname_rows(rows) -> Dict[str, str]:
    """{SIC code (string): Damodaran industry name} from a SIC mapping sheet (indname.xls layout). NOT PRESENT in the approved data."""
    hdr = None
    for i, r in enumerate(rows):
        for j, c in enumerate(r):
            if _norm(c) == "sic code":
                hdr, jc = i, j
                break
        if hdr is not None:
            break
    if hdr is None:
        raise DataNotPresent("indname sheet: no 'SIC Code' header")
    out = {}
    for r in rows[hdr + 1:]:
        if len(r) <= jc + 1 or r[jc] in (None, ""):
            continue
        code = _num(r[jc])
        if code is None:
            continue
        name = next((str(c).strip() for c in r[jc + 1:] if str(c).strip()), "")
        if name:
            out[str(int(code))] = name
    if not out:
        raise DataNotPresent("indname sheet: no rows parsed")
    return out


# ---- the store ----------------------------------------------------------------------------------------------------------------------

def default_root() -> Path:
    """DAMODARAN_DIR, else <STOCKS_ROOT>/stocks-workspace/data/damodaran. A root that does not exist raises and names the candidates."""
    import os
    import paths
    env = (os.environ.get("DAMODARAN_DIR") or "").strip()
    cands = [Path(env).expanduser()] if env else [paths.stocks_root() / "stocks-workspace" / "data" / "damodaran"]
    for c in cands:
        if (c / "MANIFEST.json").exists():
            return c
    raise DataNotPresent(f"{APPROVAL_NOTE} Tried (MANIFEST.json): {', '.join(str(c / 'MANIFEST.json') for c in cands)}")


class DamodaranStore:
    """The approved Damodaran files: `root/MANIFEST.json` lists each original .xls with its sha256; every sheet was converted to
    `root/csv/<stem>__<n>_<sheet>.csv` (stem = the .xls name without extension). The store

      * re-hashes the original .xls on every load and refuses a mismatch with the manifest;
      * reads the sheet CSVs and records BOTH the .xls sha256 and the sha256 of the CSV used in every result;
      * dates each beta file by its own 'Date updated:' cell and uses the newest dated STRICTLY BEFORE the valuation date;
      * takes the ERP from the histimpl row of the newest year-end that is at least `publication_lag_days` before the valuation date.
    """

    def __init__(self, root=None, reader: Callable = read_rows):
        self.root = Path(root) if root is not None else default_root()
        self._reader = reader
        mp = self.root / "MANIFEST.json"
        if not mp.exists():
            raise DataNotPresent(f"{APPROVAL_NOTE} No manifest at {mp}")
        self.manifest = json.loads(mp.read_text(encoding="utf-8"))
        self._cache = {}

    def _verify(self, xls_name):
        spec = (self.manifest.get("files") or {}).get(xls_name)
        p = self.root / xls_name
        if spec is None or not p.exists():
            raise DataNotPresent(f"{xls_name} is not in the manifest or not on disk at {p}")
        h = sha256_file(p)
        if h != spec["sha256"]:
            raise DataNotPresent(f"hash mismatch for {p}: file {h}, manifest {spec['sha256']} - refusing to use it")
        return h

    def _sheet_csv(self, xls_name, header_first_cell):
        stem = Path(xls_name).stem
        for c in sorted((self.root / "csv").glob(f"{stem}__*.csv")):
            rows = self._reader(c)
            if any(r and _norm(r[0]) == header_first_cell for r in rows[:30]):
                return c, rows
        raise DataNotPresent(f"no sheet of {xls_name} has a header row starting {header_first_cell!r} under {self.root / 'csv'}")

    def _xls_names(self, prefix):
        return sorted(n for n in (self.manifest.get("files") or {}) if n.startswith(prefix))

    def beta_snapshots(self) -> List[dict]:
        """[{file, date, sha256, csv, csv_sha256}] of every betas file, dated by its own 'Date updated:' cell."""
        if "snapshots" not in self._cache:
            out = []
            for n in self._xls_names("archives_betas") + self._xls_names("datasets_betas"):
                h = self._verify(n)
                c, rows = self._sheet_csv(n, "industry name")
                out.append({"file": n, "date": parse_updated_cell(rows), "sha256": h, "csv": c.name, "csv_sha256": sha256_file(c),
                            "_rows": rows})
            self._cache["snapshots"] = sorted(out, key=lambda e: e["date"])
        return self._cache["snapshots"]

    def beta_provider(self, beta_column: str = BETA_COLUMN):
        def provider(asof, industry):
            asof_d = parse_date(asof)
            ok = [e for e in self.beta_snapshots() if e["date"] < asof_d]
            if not ok:
                raise DataNotPresent(f"no betas file dated before {asof_d} (no look-ahead: a later file is not used)")
            e = ok[-1]
            table = parse_betas_rows(e["_rows"], beta_column)
            if industry not in table:
                raise DataNotPresent(f"industry {industry!r} is not in {e['file']} ({e['date']}); has {len(table)} industries")
            return (table[industry]["beta_unlevered"], f"Damodaran {e['file']} ({e['date']}), {industry}, {beta_column}",
                    {"file": e["file"], "snapshot_date": e["date"].isoformat(), "file_sha256": e["sha256"], "csv": e["csv"],
                     "csv_sha256": e["csv_sha256"]})
        return provider

    def erp_provider(self, publication_lag_days: int = PUBLICATION_LAG_DAYS):
        def provider(asof):
            asof_d = parse_date(asof)
            n = self._xls_names("datasets_histimpl")
            if not n:
                raise DataNotPresent(f"{APPROVAL_NOTE} The manifest lists no histimpl file.")
            h = self._verify(n[0])
            c, rows = self._sheet_csv(n[0], "year")
            table = parse_histimpl_rows(rows)
            ok = [y for y in table if date(y, 12, 31) + timedelta(days=publication_lag_days) < asof_d]
            if not ok:
                raise DataNotPresent(f"no year-end implied ERP at least {publication_lag_days} days before {asof_d}")
            y = max(ok)
            return (table[y], f"Damodaran implied ERP (FCFE), year-end {y}, {n[0]}",
                    {"file": n[0], "year_end": y, "file_sha256": h, "csv": c.name, "csv_sha256": sha256_file(c)})
        return provider

    def industry_for_sic(self, sic) -> str:
        """The SIC -> industry mapping (indname.xls) is NOT among the approved files: always a refusal until it is added."""
        raise DataNotPresent("the SIC -> Damodaran industry mapping (indname.xls) is not among the approved/downloaded files; "
                             "pass the industry explicitly")


# ---- the risk-free source -----------------------------------------------------------------------------------------------------------

class MriRiskFree:
    """rf(asof): the nominal 10-year yield from the MRI store on the last observation BEFORE `asof` (strictly: asof - 1 day),
    through rs2_data.asof_risk_free_legs (imported lazily: it loads the repo's config). None -> DataNotPresent; no fallback."""

    def __call__(self, asof):
        import rs2_data
        legs = rs2_data.asof_risk_free_legs((parse_date(asof) - timedelta(days=1)).isoformat())
        leg = legs.get("nominal_10y")
        if not leg:
            raise DataNotPresent(f"no nominal 10-year risk-free observation before {asof}")
        return leg[0], leg[1]


# ---- placeholders (tests and plumbing ONLY) ------------------------------------------------------------------------------------------

class PlaceholderRates:
    """Clearly labelled stand-ins so the engines can be exercised before the Damodaran files exist. Every DiscountRate built from them
    carries placeholder=True and a warning. Construction needs the explicit flag."""

    LABEL = "PLACEHOLDER_NOT_A_MEASUREMENT"

    def __init__(self, i_understand_tests_only: bool, rf=0.04, erp=0.045, beta_unlevered=1.0):
        if not i_understand_tests_only:
            raise InputError("PlaceholderRates is for tests and plumbing; pass i_understand_tests_only=True")
        self.rf_v, self.erp_v, self.bu_v = rf, erp, beta_unlevered

    def rf(self, asof):
        return self.rf_v, self.LABEL

    def beta(self, asof, industry):
        return self.bu_v, self.LABEL, {"file_sha256": None}

    def erp(self, asof):
        return self.erp_v, self.LABEL, {"file_sha256": None}
